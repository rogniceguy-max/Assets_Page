import tempfile
import unittest
from pathlib import Path
from backend.project_workbooks.nas_adapter import NasReviewServices

class TestNasReviewServices(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "reviews.db"
        self.service = NasReviewServices(self.db_path, row_kind="scene-rows")
        self.actor1 = {"user": {"id": 1, "username": "alice"}}
        self.actor2 = {"user": {"id": 2, "username": "bob"}}

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_empty_thread(self):
        res = self.service.thread(self.actor1, "file1", "row1", None)
        self.assertEqual(res["comment_count"], 0)
        self.assertEqual(len(res["comments"]), 0)

    def test_add_and_get_comment(self):
        res = self.service.add_comment(self.actor1, "file1", "row1", None, "第一条批注")
        self.assertEqual(res["comment_count"], 1)
        comment = res["comments"][0]
        self.assertEqual(comment["text"], "第一条批注")
        self.assertEqual(comment["user_name"], "alice")
        self.assertTrue(comment["can_edit"])
        self.assertTrue(comment["can_delete"])

        # Bob views thread
        res_bob = self.service.thread(self.actor2, "file1", "row1", None)
        self.assertEqual(res_bob["comment_count"], 1)
        comment_bob = res_bob["comments"][0]
        self.assertFalse(comment_bob["can_edit"])
        self.assertFalse(comment_bob["can_delete"])

    def test_edit_comment(self):
        res = self.service.add_comment(self.actor1, "file1", "row1", None, "原内容")
        c_id = res["comments"][0]["id"]
        version = res["comments"][0]["version"]

        res2 = self.service.edit_comment(self.actor1, "file1", "row1", c_id, version, "修改后内容")
        self.assertEqual(res2["comments"][0]["text"], "修改后内容")
        self.assertEqual(res2["comments"][0]["version"], version + 1)

    def test_delete_comment(self):
        res = self.service.add_comment(self.actor1, "file1", "row1", None, "待删除")
        c_id = res["comments"][0]["id"]
        version = res["comments"][0]["version"]

        res2 = self.service.delete_comment(self.actor1, "file1", "row1", c_id, version)
        self.assertEqual(res2["comment_count"], 0)

    def test_update_scene(self):
        import openpyxl
        from backend.project_workbooks.nas_adapter import NasWorkflowServices, _encode_file_id
        
        proj_dir = Path(self.temp_dir.name) / "项目"
        proj_dir.mkdir()
        wb_path = proj_dir / "场景节拍表.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "第1集"
        ws.append(["场景号", "节拍号", "故事内容", "制作难度", "拆分规划"])
        ws.append(["1", "1", "开场情节", "未标注", "原始规划"])
        wb.save(wb_path)

        workflow_svc = NasWorkflowServices(Path(self.temp_dir.name))
        file_id = _encode_file_id(str(proj_dir.stat().st_ino), "场景节拍表.xlsx")
        
        updated = workflow_svc.update_scene(self.actor1, file_id, "s0#2", 1, "红", "更新后的规划")
        self.assertEqual(updated["label_level"], "红")
        self.assertEqual(updated["breakdown_plan"], "更新后的规划")
        self.assertEqual(updated["scene_number"], "1")
        self.assertEqual(updated["beat_number"], "1")
        self.assertEqual(updated["version"], 2)

        # Re-read file to verify persistence
        wb2 = openpyxl.load_workbook(wb_path)
        ws2 = wb2.active
        self.assertEqual(ws2.cell(row=2, column=4).value, "红")
        self.assertEqual(ws2.cell(row=2, column=5).value, "更新后的规划")
