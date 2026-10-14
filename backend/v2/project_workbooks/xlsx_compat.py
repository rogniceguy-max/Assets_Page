"""Explicit XLSX compatibility boundary for migration, import, and manual exports.

The live workbook model is SQLite. This module owns the legacy spreadsheet column names and
the one-way snapshot conversion so they do not leak into the storage implementation.
"""
from __future__ import annotations

import argparse
import io
import json
from typing import Any, Mapping
from pathlib import Path

from .projects_api import ApiFault

NOTE_LABELS = (
    "成员批注一", "成员批注二", "成员批注三",
    "成员批注四", "成员批注五", "成员批注六",
)
NOTE_ALIASES = {label: label for label in NOTE_LABELS}
NOTE_ALIASES.update({f"成员批注{number}": NOTE_LABELS[number - 1]
                     for number in range(1, len(NOTE_LABELS) + 1)})
SCENE_HEADERS = ("场景号", "节拍号", "故事内容", "标注等级", "拆分规划", *NOTE_LABELS)
ASSET_HEADERS = ("类别", "名称", "出场", "形态")
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_document(document: Mapping[str, Any]) -> bytes:
    """Create an explicit XLSX snapshot from an already-read SQLite document."""
    try:
        from openpyxl import Workbook
    except ImportError:
        raise ApiFault(503, "EXPORT_UNAVAILABLE", "表格导出组件未安装。") from None

    workbook = Workbook()
    try:
        workbook.remove(workbook.active)
        for sheet_data in document.get("sheets", []):
            title = (sheet_data.get("name") or "工作表")[:31]
            if title in workbook.sheetnames:
                title = (title[:27] + "_" + str(len(workbook.sheetnames) + 1))[:31]
            sheet = workbook.create_sheet(title)
            kind = sheet_data.get("kind")
            if kind == "scene_beats":
                sheet.append(SCENE_HEADERS)
                for row in sheet_data.get("scene_rows", []):
                    notes = {note.get("label"): note.get("text")
                             for note in row.get("legacy_notes", [])}
                    sheet.append([
                        row.get("scene_number"), row.get("beat_number"), row.get("story_content"),
                        row.get("label_level"), row.get("breakdown_plan"),
                        *[notes.get(label) for label in NOTE_LABELS],
                    ])
            elif kind == "asset_review":
                sheet.append(ASSET_HEADERS)
                for row in sheet_data.get("asset_rows", []):
                    sheet.append([row.get("category"), row.get("name"),
                                  row.get("appearance"), row.get("variant")])
            else:
                columns = sheet_data.get("columns") or []
                sheet.append([column.get("label") for column in columns])
                for row in sheet_data.get("generic_rows") or []:
                    cells = row.get("cells") or {}
                    sheet.append([cells.get(column.get("key")) for column in columns])

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()


def main() -> None:
    """Manual converter for a saved API document JSON; it never opens the live database."""
    parser = argparse.ArgumentParser(description="将项目工作簿 API JSON 快照转换为 XLSX。")
    parser.add_argument("document_json", type=Path, help="GET 工作簿接口保存的 JSON 响应")
    parser.add_argument("output_xlsx", type=Path, help="输出 XLSX 路径")
    args = parser.parse_args()
    try:
        document = json.loads(args.document_json.read_text(encoding="utf-8"))
        if not isinstance(document, dict) or not isinstance(document.get("sheets"), list):
            parser.error("JSON 必须是工作簿文档对象，并包含 sheets 数组。")
        args.output_xlsx.write_bytes(export_document(document))
    except (OSError, json.JSONDecodeError, ApiFault) as error:
        parser.error(str(error))


if __name__ == "__main__":
    main()
