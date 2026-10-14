"""The xlsx projector, checked against the real 孟婆 workbooks.

These assertions are the reason the projector can be trusted at all: the contract has a dozen
ways to be subtly wrong — a dropped sheet, an invented note, a truncated body — and none of
them raise. Counting sheets and comparing text byte for byte is what makes "it read the
workbook" a claim rather than a hope.

The samples are real production files kept outside the repository, so the tests skip cleanly
when they are absent rather than failing on a machine that does not have them.
"""
from __future__ import annotations

import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

import openpyxl  # noqa: E402

from backend.v2.project_workbooks.projector import classify, project_document  # noqa: E402

SAMPLES = Path(r"C:\Users\CHEN\Desktop\temp\mengpo-workflow-audit\sources")
MASTER = SAMPLES / "孟婆2_场景节拍表.xlsx"
EPISODE_26 = SAMPLES / "孟婆2_场景节拍表_第26集.xlsx"


def _scene_rows(document):
    return [row for sheet in document["sheets"] if sheet["kind"] == "scene_beats"
            for row in sheet["scene_rows"]]


class ProjectorTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not MASTER.is_file() or not EPISODE_26.is_file():
            raise unittest.SkipTest("孟婆 samples are not present on this machine")
        cls.master = project_document(MASTER, file_id="f", name=MASTER.name, project_id="1")
        cls.episode = project_document(EPISODE_26, file_id="g", name=EPISODE_26.name, project_id="1")

    def test_every_sheet_survives(self):
        # The count is taken from openpyxl rather than written down, so adding a sheet to the
        # workbook cannot silently start failing — the assertion is "nothing was dropped".
        expected = len(openpyxl.load_workbook(MASTER, read_only=True).sheetnames)
        self.assertEqual(expected, len(self.master["sheets"]))

    def test_the_instructions_sheet_is_not_dropped(self):
        # It is neither a beat sheet nor an asset sheet, and it is the one sheet the real
        # master holds outside the per-episode pattern. Dropping it would be silent data loss.
        generic = [s for s in self.master["sheets"] if s["kind"] == "generic"]
        self.assertEqual(["使用说明"], [s["name"] for s in generic])
        self.assertTrue(generic[0]["columns"], "a generic sheet still has to carry its columns")

    def test_sheet_kinds_are_from_the_contract_vocabulary(self):
        for sheet in self.master["sheets"]:
            self.assertIn(sheet["kind"], {"asset_review", "scene_beats", "generic"})

    def test_every_beat_is_projected(self):
        self.assertEqual(209, len(_scene_rows(self.master)))

    def test_a_single_episode_file_keeps_its_own_beat_count(self):
        # The independent 第26集 file has 8 beats; the master's 第26集 sheet has 7. They are
        # different documents and reconciling them by beat number would be a fabrication.
        self.assertEqual(8, len(_scene_rows(self.episode)))
        master_26 = [s for s in self.master["sheets"] if s["name"] == "第26集"]
        self.assertEqual(7, len(master_26[0]["scene_rows"]))

    def test_six_legacy_notes_on_every_row_in_canonical_order(self):
        for row in _scene_rows(self.master):
            self.assertEqual(6, len(row["legacy_notes"]))
            self.assertEqual(
                ["成员批注一", "成员批注二", "成员批注三", "成员批注四", "成员批注五", "成员批注六"],
                [note["label"] for note in row["legacy_notes"]])

    def test_note_cells_point_at_the_real_columns(self):
        # F..K in the real master. A note column that was absent would have to carry a null
        # cell rather than a made-up address.
        first = _scene_rows(self.master)[0]
        self.assertEqual(["F2", "G2", "H2", "I2", "J2", "K2"],
                         [note["source_cell"] for note in first["legacy_notes"]])

    def test_story_text_is_byte_identical_to_the_cell(self):
        sheet = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)["第1集"]
        rows = list(sheet.iter_rows(min_row=2, max_row=2, max_col=3))
        self.assertEqual(rows[0][2].value, _scene_rows(self.master)[0]["story_content"])

    def test_difficulty_is_never_guessed(self):
        for row in _scene_rows(self.master):
            self.assertIn(row["label_level"], {"红", "黄", "绿", None})

    def test_context_is_left_null_rather_than_invented(self):
        for row in _scene_rows(self.master):
            self.assertIsNone(row["scene_context"])
            self.assertIsNone(row["cast_context"])
            self.assertEqual([], row["origins"])
            self.assertTrue(row["can_edit"])

    def test_the_document_fits_the_response_ceiling(self):
        payload = json.dumps(self.master, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        self.assertLess(len(payload), 25 * 1024 * 1024)

    def test_the_file_entry_carries_every_required_key(self):
        for key in ("id", "name", "template_kind", "version", "project_id", "revision_of", "revision_label"):
            self.assertIn(key, self.master["file"])
        self.assertEqual("scene_beats", self.master["file"]["template_kind"])

    def test_a_workbook_is_classified_by_content(self):
        self.assertEqual("scene_beats", classify(MASTER))
        self.assertEqual("scene_beats", classify(EPISODE_26))


if __name__ == "__main__":
    unittest.main()
